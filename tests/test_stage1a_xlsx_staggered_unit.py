from __future__ import annotations

from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from stroy_snab.experiments.stage1a_xlsx import (
    XlsxLineExtractionError,
    extract_xlsx_lines,
)


def test_staggered_explicit_unit_header_overrides_unlabeled_adjacent_tokens(
    tmp_path: Path,
) -> None:
    path = tmp_path / "staggered_unit_header.xlsx"
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Заявка"

    worksheet["D1"] = "Ед. изм."
    worksheet.merge_cells("D1:D2")
    worksheet["A2"] = "Наименование"
    worksheet["B2"] = "Кол-во"
    worksheet["A3"] = "Кабель"
    worksheet["B3"] = 10
    worksheet["C3"] = "м"
    worksheet["D3"] = "шт"
    worksheet["A4"] = "Провод"
    worksheet["B4"] = 20
    worksheet["C4"] = "м"
    worksheet["D4"] = "шт"

    workbook.save(path)
    workbook.close()

    lines = extract_xlsx_lines(
        path,
        document_id="REQUEST_0025",
        document_role="REQUEST",
    )

    assert len(lines) == 2
    assert [line.quantity for line in lines] == [Decimal("10"), Decimal("20")]
    assert [line.unit_raw for line in lines] == ["шт", "шт"]


def test_multiple_explicit_unit_headers_in_contiguous_header_block_fail_closed(
    tmp_path: Path,
) -> None:
    path = tmp_path / "ambiguous_staggered_units.xlsx"
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Заявка"

    worksheet["D1"] = "Ед. изм."
    worksheet["A2"] = "Наименование"
    worksheet["B2"] = "Кол-во"
    worksheet["E2"] = "Единица измерения"
    worksheet["A3"] = "Кабель"
    worksheet["B3"] = 10
    worksheet["C3"] = "м"
    worksheet["D3"] = "шт"
    worksheet["E3"] = "м"

    workbook.save(path)
    workbook.close()

    with pytest.raises(
        XlsxLineExtractionError,
        match="ambiguous explicit unit headers",
    ):
        extract_xlsx_lines(
            path,
            document_id="REQUEST_0026",
            document_role="REQUEST",
        )


def test_staggered_explicit_unit_collision_with_item_fails_closed(
    tmp_path: Path,
) -> None:
    path = tmp_path / "unit_item_collision.xlsx"
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Заявка"

    worksheet["A1"] = "Ед. изм."
    worksheet["A2"] = "Наименование"
    worksheet["B2"] = "Кол-во"
    worksheet["A3"] = "Кабель"
    worksheet["B3"] = 10
    worksheet["C3"] = "м"
    worksheet["A4"] = "Провод"
    worksheet["B4"] = 20
    worksheet["C4"] = "шт"

    workbook.save(path)
    workbook.close()

    with pytest.raises(
        XlsxLineExtractionError,
        match="explicit unit header collides",
    ):
        extract_xlsx_lines(
            path,
            document_id="REQUEST_0027",
            document_role="REQUEST",
        )


def test_vertical_merged_non_unit_header_blocks_adjacent_unit_inference(
    tmp_path: Path,
) -> None:
    path = tmp_path / "merged_comment_header.xlsx"
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Заявка"

    worksheet["C1"] = "Комментарий"
    worksheet.merge_cells("C1:C2")
    worksheet["A2"] = "Наименование"
    worksheet["B2"] = "Кол-во"
    worksheet["A3"] = "Кабель"
    worksheet["B3"] = 10
    worksheet["C3"] = "м"
    worksheet["A4"] = "Провод"
    worksheet["B4"] = 20
    worksheet["C4"] = "шт"

    workbook.save(path)
    workbook.close()

    lines = extract_xlsx_lines(
        path,
        document_id="REQUEST_0028",
        document_role="REQUEST",
    )

    assert [line.unit_raw for line in lines] == [None, None]


def test_horizontal_merged_non_unit_header_blocks_adjacent_unit_inference(
    tmp_path: Path,
) -> None:
    path = tmp_path / "horizontal_merged_comment_header.xlsx"
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Заявка"

    worksheet["B1"] = "Комментарий"
    worksheet.merge_cells("B1:C1")
    worksheet["A2"] = "Наименование"
    worksheet["B2"] = "Кол-во"
    worksheet["A3"] = "Кабель"
    worksheet["B3"] = 10
    worksheet["C3"] = "м"
    worksheet["A4"] = "Провод"
    worksheet["B4"] = 20
    worksheet["C4"] = "шт"

    workbook.save(path)
    workbook.close()

    lines = extract_xlsx_lines(
        path,
        document_id="REQUEST_0029",
        document_role="REQUEST",
    )

    assert [line.unit_raw for line in lines] == [None, None]
