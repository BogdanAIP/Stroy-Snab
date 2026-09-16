from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
import re

import openpyxl


class XlsxLineExtractionError(RuntimeError):
    """Raised when the Stage 1A XLSX baseline cannot safely extract a supported table."""


@dataclass(frozen=True, slots=True)
class ProcurementLine:
    document_id: str
    document_role: str
    item_name_raw: str
    unit_raw: str | None
    quantity: Decimal
    source_locator: str


_HEADER_SCAN_LIMIT = 60
_TOTAL_PREFIXES = ("итого", "всего")
_QUANTITY_HEADERS = {
    "кол во",
    "количество",
    "количество объем",
    "объем",
}
_STRICT_NUMBER = r"[+-]?(?:\d{1,3}(?:\s\d{3})+|\d+)(?:[.,]\d+)?"
_UNIT_SUFFIX = r"[A-Za-zА-Яа-яЁё%].*"


def _normalize_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\u00a0", " ").replace("ё", "е").strip().lower()
    text = re.sub(r"[\\/_.()\-]+", " ", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def _header_role(value: object) -> str | None:
    text = _normalize_text(value)
    if not text:
        return None

    if "наименован" in text or text in {"товар", "материал", "товары работы услуги"}:
        return "item"

    if text in _QUANTITY_HEADERS:
        return "quantity"

    if "единица измерения" in text:
        return "unit"
    if text in {"ед", "ед изм", "ед измер", "единица"}:
        return "unit"
    if text.startswith("ед ") and "изм" in text:
        return "unit"

    return None


def _as_quantity(value: object) -> tuple[Decimal | None, str | None]:
    if value is None or isinstance(value, bool):
        return None, None
    if isinstance(value, Decimal):
        return value, None
    if isinstance(value, (int, float)):
        return Decimal(str(value)), None

    text = str(value).replace("\u00a0", " ").strip()
    if not text:
        return None, None

    match = re.fullmatch(
        rf"({_STRICT_NUMBER})(?:\s+({_UNIT_SUFFIX}))?",
        text,
    )
    if match is None:
        return None, None

    number_text = re.sub(r"\s+", "", match.group(1)).replace(",", ".")
    try:
        quantity = Decimal(number_text)
    except InvalidOperation:
        return None, None

    unit_hint = match.group(2)
    if unit_hint is not None:
        unit_hint = unit_hint.strip() or None
    return quantity, unit_hint


def _find_header_columns(worksheet) -> tuple[int, dict[str, int]] | None:
    max_row = min(worksheet.max_row or _HEADER_SCAN_LIMIT, _HEADER_SCAN_LIMIT)
    for row_index, row in enumerate(
        worksheet.iter_rows(min_row=1, max_row=max_row),
        start=1,
    ):
        candidates: dict[str, list[int]] = {}
        for column_index, cell in enumerate(row, start=1):
            role = _header_role(cell.value)
            if role is not None:
                candidates.setdefault(role, []).append(column_index)

        if "item" not in candidates or "quantity" not in candidates:
            continue

        ambiguous = {
            role: columns
            for role, columns in candidates.items()
            if len(columns) > 1
        }
        if ambiguous:
            roles = ", ".join(sorted(ambiguous))
            raise XlsxLineExtractionError(
                f"ambiguous procurement table header roles: {roles}"
            )

        return row_index, {
            role: columns[0]
            for role, columns in candidates.items()
        }

    return None


def _is_formula(cell) -> bool:
    return cell.data_type == "f" or (
        isinstance(cell.value, str) and cell.value.startswith("=")
    )


def extract_xlsx_lines(
    path: str | Path,
    *,
    document_id: str,
    document_role: str,
) -> tuple[ProcurementLine, ...]:
    """Extract procurement lines from a native XLSX using a deterministic baseline.

    This is a Stage 1A experiment, not a general spreadsheet parser. It looks for
    a table with item-name and quantity headers, optionally a unit header, and
    fails closed when a candidate row cannot be interpreted safely.
    """

    workbook = openpyxl.load_workbook(
        filename=Path(path),
        read_only=True,
        data_only=False,
    )

    extracted: list[ProcurementLine] = []
    detected_table = False
    try:
        for worksheet in workbook.worksheets:
            header = _find_header_columns(worksheet)
            if header is None:
                continue

            detected_table = True
            header_row, roles = header
            for row in worksheet.iter_rows(min_row=header_row + 1):
                item_cell = row[roles["item"] - 1]
                item_name = "" if item_cell.value is None else str(item_cell.value).strip()
                if not item_name:
                    continue

                normalized_item = _normalize_text(item_name)
                if normalized_item.startswith(_TOTAL_PREFIXES):
                    break

                quantity_cell = row[roles["quantity"] - 1]
                if _is_formula(quantity_cell):
                    raise XlsxLineExtractionError(
                        f"formula quantity is unsupported at "
                        f"{worksheet.title}!{quantity_cell.coordinate}"
                    )

                quantity, quantity_unit = _as_quantity(quantity_cell.value)
                if quantity is None:
                    raise XlsxLineExtractionError(
                        f"unparseable quantity at "
                        f"{worksheet.title}!{quantity_cell.coordinate}"
                    )

                unit_raw: str | None = quantity_unit
                unit_column = roles.get("unit")
                if unit_column is not None:
                    unit_cell = row[unit_column - 1]
                    if _is_formula(unit_cell):
                        raise XlsxLineExtractionError(
                            f"formula unit is unsupported at "
                            f"{worksheet.title}!{unit_cell.coordinate}"
                        )
                    if unit_cell.value is not None:
                        value = str(unit_cell.value).strip()
                        if value:
                            unit_raw = value

                extracted.append(
                    ProcurementLine(
                        document_id=document_id,
                        document_role=document_role,
                        item_name_raw=item_name,
                        unit_raw=unit_raw,
                        quantity=quantity,
                        source_locator=f"{worksheet.title}!{item_cell.coordinate}",
                    )
                )
    finally:
        workbook.close()

    if not detected_table:
        raise XlsxLineExtractionError("no supported procurement table header found")
    if not extracted:
        raise XlsxLineExtractionError("supported table found but no procurement lines extracted")

    return tuple(extracted)
